from __future__ import annotations

from pathlib import Path
import asyncio
import time

import pytest
from docx import Document
from pptx import Presentation
from fastapi.testclient import TestClient

from app.core.config import Settings
from app.models.goal_research_schemas import (
    CriterionEvaluation, CriterionStatus, ExpectedResultStatus,
    GoalCriterion, GoalResearchResponse, GoalRunStatus, GoalStatus, GoalStopReason,
    GoalResearchRequest,
)
from app.services.deliverables.artifact_validator import validate_artifact
from app.services.deliverables.service import DeliverableService
from app.services.goal_research_service import GoalResearchRunNotFound, GoalResearchService
from app.api.goal_research_routes import get_goal_research_service
from app.main import app


def frozen_result():
    return GoalResearchResponse(
        run_id="GR-TEST", topic="RFT pressure comparison", goal="Compare conditions",
        expected_result="Pressure always increases", run_status=GoalRunStatus.COMPLETED,
        status=GoalStatus.ACHIEVED, stop_reason=GoalStopReason.GOAL_ACHIEVED,
        expected_result_status=ExpectedResultStatus.CONTRADICTED,
        final_answer="The hypothesis is contradicted by the observed evidence [KB1].",
        max_iterations=2, iterations_completed=2,
        frozen_criteria=[GoalCriterion(criterion_id="C1", description="Compare conditions")],
        criteria=[CriterionEvaluation(criterion_id="C1", status=CriterionStatus.MET, reason="Evidence", supporting_evidence=["KB1"])],
    )


def test_docx_generation_and_reopen(tmp_path: Path):
    settings = Settings(data_dir=tmp_path / "data", agent_workspace_dir=tmp_path / "workspace")
    artifact = DeliverableService(settings).create(frozen_result(), "docx")
    path = settings.agent_workspace_dir / artifact.path
    validate_artifact(path, "docx")
    document = Document(path)
    all_text = "\n".join(item.text for item in document.paragraphs)
    assert "RFT pressure comparison" in all_text
    assert "Compare conditions" in all_text
    assert "contradicted" in all_text
    assert "supported" not in all_text.lower().replace("not supported", "")
    assert len(document.tables) >= 3 and artifact.size_bytes > 2000


def test_pptx_generation_and_reopen(tmp_path: Path):
    settings = Settings(data_dir=tmp_path / "data", agent_workspace_dir=tmp_path / "workspace")
    artifact = DeliverableService(settings).create(frozen_result(), "pptx")
    path = settings.agent_workspace_dir / artifact.path
    validate_artifact(path, "pptx")
    deck = Presentation(path)
    text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert len(deck.slides) >= 5
    assert "Main Results" in text and "Limitations" in text and "Conclusion" in text
    assert "contradicted" in text
    assert artifact.size_bytes > 2000


def test_deliverable_does_not_mutate_frozen_research_and_avoids_overwrite(tmp_path: Path):
    settings = Settings(data_dir=tmp_path / "data", agent_workspace_dir=tmp_path / "workspace")
    result = frozen_result()
    before = result.model_dump(mode="json")
    service = DeliverableService(settings)
    first = service.create(result, "docx")
    second = service.create(result, "docx")
    assert first.path != second.path
    assert result.model_dump(mode="json") == before


def test_download_path_is_whitelisted_to_run_deliverables(tmp_path: Path):
    settings = Settings(data_dir=tmp_path / "data", agent_workspace_dir=tmp_path / "workspace")
    artifact = DeliverableService(settings).create(frozen_result(), "docx")
    result = frozen_result().model_copy(update={"artifacts": [artifact]})
    service = GoalResearchService(settings, object())
    service._write("GR-TEST", GoalResearchRequest(topic="x"), result)
    assert service.artifact_path("GR-TEST", artifact.artifact_id).is_file()
    with pytest.raises(GoalResearchRunNotFound):
        service.artifact_path("GR-TEST", "../secrets")
    app.dependency_overrides[get_goal_research_service] = lambda: service
    try:
        with TestClient(app) as client:
            valid = client.get("/api/research/goal-runs/GR-TEST/artifacts/ART-DOCX")
            assert valid.status_code == 200
            assert valid.content[:2] == b"PK"
            assert client.get("/api/research/goal-runs/GR-TEST/artifacts/ART-UNKNOWN").status_code == 404
    finally:
        app.dependency_overrides.clear()


def test_deliverables_start_only_after_research_finishes(tmp_path: Path):
    settings = Settings(data_dir=tmp_path / "data", agent_workspace_dir=tmp_path / "workspace")
    output_dir = settings.agent_workspace_dir / "results" / "goal-research"

    class Controller:
        async def run(self, run_id, request, *, on_progress, is_canceled):
            running = GoalResearchResponse(run_id=run_id, topic=request.topic, run_status=GoalRunStatus.RUNNING, max_iterations=1)
            on_progress(running)
            assert not (output_dir / run_id / "deliverables").exists()
            await asyncio.sleep(0.02)
            result = frozen_result()
            result.run_id = run_id
            return result

    service = GoalResearchService(settings, Controller())
    started = service.start(GoalResearchRequest(topic="x", deliverables=["docx", "pptx"], max_iterations=1))
    deadline = time.monotonic() + 5
    response = started
    while time.monotonic() < deadline:
        response = service.get(started.run_id)
        if response.deliverable_status.get("docx") == "completed" and response.deliverable_status.get("pptx") == "completed" and "docx_success" in response.telemetry:
            break
        time.sleep(0.02)
    assert response.status == GoalStatus.ACHIEVED
    assert response.deliverable_status == {"docx": "completed", "pptx": "completed"}
    assert len(response.artifacts) == 2
    assert response.final_answer == frozen_result().final_answer
    assert response.telemetry["docx_success"] and response.telemetry["pptx_success"]


def test_failed_docx_does_not_change_research_or_block_pptx(tmp_path: Path):
    settings = Settings(data_dir=tmp_path / "data", agent_workspace_dir=tmp_path / "workspace")

    class Controller:
        async def run(self, run_id, request, *, on_progress, is_canceled):
            result = frozen_result()
            result.run_id = run_id
            return result

    service = GoalResearchService(settings, Controller())
    original = service.deliverables.create

    def create(frozen, kind, include_charts=True):
        if kind == "docx":
            raise ValueError("bad DOCX")
        return original(frozen, kind, include_charts)

    service.deliverables.create = create
    started = service.start(GoalResearchRequest(topic="x", deliverables=["docx", "pptx"]))
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        response = service.get(started.run_id)
        if response.deliverable_status.get("pptx") == "completed":
            break
        time.sleep(0.02)
    assert response.status == GoalStatus.ACHIEVED
    assert response.deliverable_status == {"docx": "failed", "pptx": "completed"}
    assert len(response.artifacts) == 1
