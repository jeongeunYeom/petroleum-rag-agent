from __future__ import annotations

import asyncio
import json
import time
from pathlib import Path

from docx import Document
from pptx import Presentation

from app.core.config import Settings
from app.models.goal_research_schemas import (
    ComputationRecord, CriterionEvaluation, CriterionStatus, ExpectedResultStatus, GoalCriterion,
    GoalIterationRecord, GoalResearchPlan, GoalResearchRequest, GoalResearchResponse,
    GoalRunStatus, GoalStatus, GoalStopReason,
)
from app.models.research_schemas import InternalEvidence
from app.services.deliverables.service import DeliverableService
from app.services.goal_evaluator import GoalEvaluator
from app.services.goal_research_agent import GoalResearchAgent
from app.services.goal_result_summary import resolve_final_limitations


class InventedHypothesisOllama:
    async def chat_structured(self, _messages, schema, **_kwargs):
        if "claims" in schema["properties"]:
            return json.dumps({
                "claims": [
                    {"claim": "Case B is 20 mD.", "citations": ["KB1"]},
                    {"claim": "The hypothesis is supported.", "citations": ["KB1"]},
                ],
                "hypothesis_assessment": {"claim": "The hypothesis is supported.", "citations": ["KB1"]},
                "limitations": [],
            })
        return json.dumps({
            "criteria": [{"criterion_id": "C1", "status": "met", "reason": "observed", "supporting_evidence": ["KB1"]}],
            "expected_result_status": "supported",
            "gaps": [],
            "next_research_need": None,
            "goal_conflicts_with_evidence": False,
        })


def _iteration(number: int, status: CriterionStatus, gap: str) -> GoalIterationRecord:
    return GoalIterationRecord(
        iteration=number, research_query="case B", plan=GoalResearchPlan(
            research_question="case B", focus_criteria=["C1"], reason="verify value",
        ), candidate_answer="Case B is 20 mD [KB2]" if number > 1 else "Case B unknown",
        criteria=[CriterionEvaluation(
            criterion_id="C1", status=status, reason="checked",
            supporting_evidence=["KB2"] if status == CriterionStatus.MET else [],
        )], goal_coverage=1 if status == CriterionStatus.MET else 0,
        expected_result_status=ExpectedResultStatus.NOT_PROVIDED,
        engineering_validation_passed=True, gap_analysis=[gap] if gap else [],
    )


def _result(last_gap: str = "Case B value is missing") -> GoalResearchResponse:
    return GoalResearchResponse(
        run_id="GR-CONSISTENCY", topic="Compare permeability", goal="Compare case A and B",
        run_status=GoalRunStatus.COMPLETED, status=GoalStatus.ACHIEVED,
        stop_reason=GoalStopReason.GOAL_ACHIEVED, max_iterations=2,
        iterations_completed=2, goal_coverage=1, goal_coverage_percent=100,
        expected_result_status=ExpectedResultStatus.NOT_PROVIDED,
        frozen_criteria=[GoalCriterion(criterion_id="C1", description="Identify both source-supported permeability inputs")],
        criteria=[CriterionEvaluation(criterion_id="C1", status=CriterionStatus.MET, reason="Both found", supporting_evidence=["KB1", "KB2"])],
        iterations=[_iteration(1, CriterionStatus.UNMET, "Case B value is missing"), _iteration(2, CriterionStatus.MET, last_gap)],
        internal_sources=[
            InternalEvidence(evidence_id="KB1", document="fixture", page=1, chunk_id="a", score=1, excerpt="Measured permeability for case A is 10 mD."),
            InternalEvidence(evidence_id="KB2", document="fixture", page=2, chunk_id="b", score=1, excerpt="Measured permeability for case B is 20 mD."),
        ],
        final_answer="Case A is 10 mD [KB1]. Case B is 20 mD [KB2].\n\nLimitations: Case B value is missing",
    )


def test_missing_hypothesis_is_ignored_in_synthesis_and_evaluation():
    llm = InventedHypothesisOllama()
    request = GoalResearchRequest(topic="Compare permeability", expected_result="   ", engineering_validation=False)
    evidence = [{"evidence_id": "KB1", "source_type": "knowledge_base", "locator": "fixture", "text": "Case B is 20 mD."}]
    candidate = asyncio.run(GoalResearchAgent(object(), llm)._synthesize(request, [], evidence, []))
    evaluation = asyncio.run(GoalEvaluator(llm).evaluate(
        request, [GoalCriterion(criterion_id="C1", description="Find case B")], candidate, evidence, {},
    ))
    assert request.expected_result is None
    assert "hypothesis is supported" not in candidate.lower()
    assert evaluation.expected_result_status == ExpectedResultStatus.NOT_PROVIDED


def test_internal_diagnostic_code_is_not_shown_as_a_user_limitation():
    class DiagnosticOllama:
        async def chat_structured(self, *_args, **_kwargs):
            return json.dumps({
                "claims": [{"claim": "The source states that case B is 20 mD.", "citations": ["KB1"]}],
                "hypothesis_assessment": {"claim": "", "citations": []},
                "limitations": ["provenance_attribution_error"],
            })

    request = GoalResearchRequest(topic="Case B permeability", engineering_validation=False)
    evidence = [{"evidence_id": "KB1", "source_type": "knowledge_base", "locator": "fixture",
                 "text": "The source states that case B is 20 mD."}]
    answer = asyncio.run(GoalResearchAgent(object(), DiagnosticOllama())._synthesize(
        request, [], evidence, []))
    assert "provenance_attribution_error" not in answer


def test_resolved_gap_is_removed_but_iteration_history_is_preserved():
    result = _result()
    result.expected_result_status = ExpectedResultStatus.SUPPORTED
    assert resolve_final_limitations(result) == []
    GoalResearchAgent._finish(result, GoalRunStatus.COMPLETED, GoalStatus.ACHIEVED, GoalStopReason.GOAL_ACHIEVED, result.final_answer, time.perf_counter())
    assert result.final_limitations == []
    assert "Case B value is missing" not in result.final_answer
    assert result.iterations[0].gap_analysis == ["Case B value is missing"]
    assert result.iterations[1].gap_analysis == ["Case B value is missing"]
    assert result.expected_result_status == ExpectedResultStatus.NOT_PROVIDED
    result.iterations[-1].gap_analysis = ["Insufficient evidence for case B permeability value in some prior findings."]
    assert resolve_final_limitations(result) == []
    result.final_answer += "\n\nLimitations: The prior findings indicated insufficient evidence for case B permeability, but this has been resolved with KB2."
    assert resolve_final_limitations(result) == []
    result.final_answer += "\n\nLimitations: The chart cannot be generated without the permeability value for case B."
    assert resolve_final_limitations(result) == []


def test_genuine_final_limitation_survives_achieved_goal():
    result = _result("Long-term validation is unavailable")
    result.final_answer = "Case B is 20 mD [KB2]."
    result.criteria.append(CriterionEvaluation(criterion_id="C2", status=CriterionStatus.PARTIAL, reason="No field study"))
    assert resolve_final_limitations(result) == ["Long-term validation is unavailable"]


def test_chart_limitation_requires_a_validated_chart_and_met_criterion():
    result = _result("")
    result.final_answer = "Case B is 20 mD [KB2].\n\nLimitations: The chart cannot be generated without both measured values."
    assert resolve_final_limitations(result) == ["The chart cannot be generated without both measured values."]
    result.frozen_criteria.append(GoalCriterion(criterion_id="C3", description="Generate a chart of both measured values"))
    result.criteria.append(CriterionEvaluation(criterion_id="C3", status=CriterionStatus.MET, reason="chart validated", supporting_evidence=["CALC1"]))
    result.computations.append(ComputationRecord(
        computation_id="CALC1", analysis_id="PA-1", purpose="chart",
        source_evidence_ids=["KB1", "KB2"], output_files=["analysis_001.png"],
        validation_passed=True,
    ))
    assert resolve_final_limitations(result) == []


def test_docx_and_pptx_share_final_state_and_keep_hypothesis_not_applicable(tmp_path: Path):
    result = _result()
    GoalResearchAgent._finish(result, GoalRunStatus.COMPLETED, GoalStatus.ACHIEVED, GoalStopReason.GOAL_ACHIEVED, result.final_answer, time.perf_counter())
    settings = Settings(data_dir=tmp_path / "data", agent_workspace_dir=tmp_path / "workspace")
    service = DeliverableService(settings)
    docx = service.create(result, "docx")
    pptx = service.create(result, "pptx")
    doc = Document(settings.agent_workspace_dir / docx.path)
    deck = Presentation(settings.agent_workspace_dir / pptx.path)
    doc_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    slides = {
        slide.shapes[0].text: "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        for slide in deck.slides
    }
    assert "No hypothesis provided." in doc_text
    assert "Not applicable" in doc_text
    assert "Hypothesis: Not provided" in slides["Goal and Hypothesis"]
    assert "Assessment: Not applicable" in slides["Hypothesis Evaluation"]
    assert "supported" not in slides["Hypothesis Evaluation"].lower()
    assert "Iteration 1 historical gap: Case B value is missing" in doc_text
    limitation_index = next(index for index, paragraph in enumerate(doc.paragraphs) if paragraph.text == "11 Limitations")
    assert "Case B value is missing" not in doc.paragraphs[limitation_index + 1].text
    assert "Case B value is missing" not in slides["Limitations"]
    assert "No unresolved limitation" in doc_text and "No unresolved limitation" in slides["Limitations"]
    assert "Goal status: achieved" in doc_text and "Goal status: achieved" in slides["Research Background"]
    assert "Coverage: 100%" in doc_text and "Coverage: 100%" in slides["Main Results"]
