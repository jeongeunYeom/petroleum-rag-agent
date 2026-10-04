"""Secondary, non-ranking DOCX/PPTX reliability check on three frozen C outputs."""

from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts.run_agentic_heldout import MANIFEST, _save


SELECTED = {"AG-WT-001": "conceptual", "AG-Q-009": "quantitative", "AG-IE-014": "insufficient_evidence"}


def _docx_text(path: Path) -> tuple[str, int]:
    from docx import Document

    doc = Document(path)
    text = "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    return text, len(doc.inline_shapes)


def _pptx_text(path: Path) -> tuple[str, int]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(path)
    text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for shape in slide.shapes)
    return text, pictures


def check_one(response, artifact, kind: str, workspace: Path) -> dict:
    path = workspace / artifact.path
    text, pictures = _docx_text(path) if kind == "docx" else _pptx_text(path)
    expected_status = response.status.value
    expected_hypothesis = response.expected_result_status.value if response.expected_result else "Not applicable"
    source_ids = [s.evidence_id for s in response.internal_sources]
    source_ids += [s.evidence_id for s in response.figures]
    valid_computations = [c for c in response.computations if c.validation_passed]
    chart_paths = [p for c in valid_computations for p in c.output_files if p.lower().endswith(".png") and (workspace / p).is_file()]
    checks = {
        "artifact_generation_success": bool(artifact.validation_passed and path.is_file() and artifact.size_bytes > 2000),
        "state_consistency": f"Goal status: {expected_status}" in text or f"Goal: {expected_status}" in text,
        "hypothesis_consistency": expected_hypothesis in text,
        "provenance_presence": "Source:" in text and (not source_ids or any(s in text for s in source_ids)),
        "calc_provenance_presence": not valid_computations or all(c.computation_id in text for c in valid_computations),
        "limitation_consistency": not response.final_limitations or all(limitation[:80] in text for limitation in response.final_limitations),
        "chart_inclusion_if_applicable": not chart_paths or pictures > 0,
    }
    return {"artifact_type": kind, "path": str(path), "size_bytes": path.stat().st_size, "picture_count": pictures, "source_chart_count": len(chart_paths), "checks": checks, "passed": all(checks.values())}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir", type=Path, required=True)
    args = parser.parse_args()
    output = args.run_dir / "secondary_artifacts.json"
    if output.exists():
        raise FileExistsError(output)
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    os.environ.setdefault("DATA_DIR", manifest["kb"]["data_dir"])
    os.environ.setdefault("RETRIEVAL_MODE", manifest["kb"]["retrieval_mode"])
    from app.core.config import get_settings
    from app.models.goal_research_schemas import GoalResearchResponse
    from app.services.deliverables.service import DeliverableService

    payload = json.loads((args.run_dir / "full_agent.json").read_text(encoding="utf-8"))
    if not payload["complete"]:
        raise ValueError("Full Agent run is not complete")
    rows = {row["task_id"]: row for row in payload["results"]}
    service = DeliverableService(get_settings())
    workspace = get_settings().agent_workspace_dir
    results = []
    for task_id, category in SELECTED.items():
        response = GoalResearchResponse.model_validate(rows[task_id]["response"])
        for kind in ("docx", "pptx"):
            try:
                artifact = service.create(response.model_copy(deep=True), kind, include_charts=True)
                result = check_one(response, artifact, kind, workspace)
            except Exception as exc:
                result = {"artifact_type": kind, "passed": False, "error": f"{type(exc).__name__}: {exc}"[:500]}
            results.append({"task_id": task_id, "category": category, **result})
            print(f"{task_id} {kind} passed={result['passed']}", flush=True)
    _save(output, {"run_id": args.run_dir.name, "core_score_included": False, "results": results})
    print(f"output={output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
