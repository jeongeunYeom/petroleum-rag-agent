from __future__ import annotations

from pathlib import Path
from zipfile import is_zipfile

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def validate_artifact(path: Path, kind: str) -> None:
    if not path.is_file() or path.stat().st_size < 2000 or not is_zipfile(path):
        raise ValueError(f"Invalid {kind} container")
    if kind == "docx":
        document = Document(path)
        headings = {paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")}
        required = {"2 Research Goal", "8 Main Findings", "10 Goal Criteria Evaluation", "13 Sources and Provenance"}
        if not required.issubset(headings) or not document.tables:
            raise ValueError("DOCX required sections or tables are missing")
        if not any(paragraph.text.strip() for paragraph in document.paragraphs):
            raise ValueError("DOCX has no text")
        for relation in document.part.rels.values():
            if relation.reltype.endswith("/image") and not relation.target_part.blob:
                raise ValueError("DOCX contains an empty image relation")
    elif kind == "pptx":
        deck = Presentation(path)
        if len(deck.slides) < 5:
            raise ValueError("PPTX has too few slides")
        slide_text = [" ".join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in deck.slides]
        if any(not value.strip() for value in slide_text) or not any("Main Results" in value for value in slide_text) or not any("Conclusion" in value for value in slide_text):
            raise ValueError("PPTX has empty or missing sections")
        for slide in deck.slides:
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and not shape.image.blob:
                    raise ValueError("PPTX contains an empty image relation")
    else:
        raise ValueError("Unsupported artifact kind")
