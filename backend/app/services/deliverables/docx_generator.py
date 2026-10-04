from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.models.goal_research_schemas import GoalResearchResponse
from app.services.deliverables.common import metadata_lines, provenance, safe_text


def _table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = safe_text(value)


def generate_docx(result: GoalResearchResponse, path: Path, chart_paths: list[Path]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.7)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.name = "Aptos"
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    document.add_paragraph(safe_text(result.topic), style="Title")
    document.add_paragraph("Goal Research Report")
    for line in metadata_lines(result):
        document.add_paragraph(safe_text(line))
    document.add_page_break()

    def section_text(heading: str, text: object) -> None:
        document.add_heading(heading, level=1)
        document.add_paragraph(safe_text(text) or "Not available.")

    section_text("1 Research Topic", result.topic)
    section_text("2 Research Goal", result.goal or result.topic)
    section_text("3 Expected Hypothesis", result.expected_result or "No hypothesis provided.")
    section_text("4 Research Method", "Iterative evidence retrieval, synthesis, criterion evaluation and bounded replanning. Validated local calculations are identified separately as CALC records.")
    document.add_heading("5 Evidence Used", level=1)
    evidence_rows = []
    for item in result.internal_sources:
        evidence_rows.append([item.evidence_id, "KB", f"{Path(item.document).name} p.{item.page or '?'}"])
    for item in result.web_sources:
        evidence_rows.append([item.evidence_id, "WEB", safe_text(item.url)])
    for item in result.figures:
        evidence_rows.append([item.evidence_id, "FIG", f"{Path(item.document).name} p.{item.page or '?'}"])
    _table(document, ["ID", "Type", "Source"], evidence_rows or [["None", "—", "No evidence recorded"]])
    document.add_heading("6 Iterative Research Process", level=1)
    _table(document, ["Iteration", "Coverage", "Query", "Calculations"], [
        [str(item.iteration), f"{item.goal_coverage:.0%}", item.research_query, ", ".join(item.computation_ids) or "—"]
        for item in result.iterations
    ] or [["—", "—", "No completed iteration", "—"]])
    document.add_heading("7 Calculation and Data Analysis", level=1)
    _table(document, ["ID", "Purpose", "Sources", "Status"], [
        [item.computation_id, item.purpose, ", ".join(item.source_evidence_ids), "validated" if item.validation_passed else "failed"]
        for item in result.computations
    ] or [["—", "No calculations", "—", "—"]])
    for chart in chart_paths:
        document.add_picture(str(chart), width=Inches(5.8))
        caption = document.add_paragraph(f"Validated analysis chart {chart.stem}")
        caption.style = "Caption"
    section_text("8 Main Findings", result.final_answer or "No supported finding was established.")
    section_text("9 Hypothesis Evaluation", f"{result.expected_result_status.value}. {result.expected_result or 'No hypothesis provided.'}")
    document.add_heading("10 Goal Criteria Evaluation", level=1)
    by_id = {item.criterion_id: item for item in result.criteria}
    _table(document, ["ID", "Criterion", "Status", "Evidence"], [
        [item.criterion_id, item.description, by_id[item.criterion_id].status.value if item.criterion_id in by_id else "unmet", ", ".join(by_id[item.criterion_id].supporting_evidence) if item.criterion_id in by_id else "—"]
        for item in result.frozen_criteria
    ] or [["—", "No criteria", "—", "—"]])
    gaps = [gap for item in result.iterations for gap in item.gap_analysis]
    section_text("11 Limitations", "; ".join(dict.fromkeys(gaps)) or "No specific gap was recorded.")
    section_text("12 Conclusion", f"Goal status: {result.status.value}. Stop reason: {result.stop_reason.value if result.stop_reason else 'not recorded'}. {result.final_answer}")
    section_text("13 Sources and Provenance", provenance(result))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.save(path)
