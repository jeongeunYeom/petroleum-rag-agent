from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from PIL import Image

from app.models.goal_research_schemas import GoalResearchResponse
from app.services.deliverables.common import metadata_lines, provenance, safe_text


def generate_pptx(result: GoalResearchResponse, path: Path, chart_paths: list[Path]) -> None:
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)

    def slide(title: str, body: list[str], source: str | None = None) -> None:
        page = deck.slides.add_slide(deck.slide_layouts[6])
        heading = page.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1)).text_frame
        heading.text = safe_text(title)
        heading.paragraphs[0].font.size = Pt(29 if len(title) > 60 else 32)
        heading.paragraphs[0].font.bold = True
        heading.paragraphs[0].font.color.rgb = RGBColor(20, 35, 55)
        text = page.shapes.add_textbox(Inches(0.95), Inches(1.75), Inches(11.4), Inches(4.9)).text_frame
        text.word_wrap = True
        for index, value in enumerate(body[:6]):
            paragraph = text.paragraphs[0] if index == 0 else text.add_paragraph()
            paragraph.text = safe_text(value)[:220]
            paragraph.font.size = Pt(21)
            paragraph.space_after = Pt(17)
        foot = page.shapes.add_textbox(Inches(0.9), Inches(7.0), Inches(11.5), Inches(0.3)).text_frame
        foot.text = safe_text(source or provenance(result))[:220]
        foot.paragraphs[0].font.size = Pt(10)
        foot.paragraphs[0].font.color.rgb = RGBColor(90, 100, 110)

    slide(safe_text(result.topic)[:90], ["Goal Research Presentation", *metadata_lines(result)[:5]])
    slide("Research Background", [result.topic, f"Goal status: {result.status.value}"])
    slide("Goal and Hypothesis", [result.goal or result.topic, f"Hypothesis: {result.expected_result or 'Not provided'}", f"Assessment: {result.expected_result_status.value if result.expected_result else 'Not applicable'}"])
    slide("Research Method", ["Iterative retrieval, evidence accumulation and synthesis", "Criterion evaluation and gap-directed replanning", "Approved, validated calculations are recorded as CALC results"])
    slide("Evidence", [
        f"Knowledge base sources: {len(result.internal_sources)}",
        f"Web sources: {len(result.web_sources)}",
        f"Figure sources: {len(result.figures)}",
        provenance(result),
    ])
    if chart_paths:
        for index, chart in enumerate(chart_paths[:2]):
            page = deck.slides.add_slide(deck.slide_layouts[6])
            title = page.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(0.7)).text_frame
            title.text = "Quantitative Analysis" if index == 0 else "Additional Analysis Chart"
            title.paragraphs[0].font.size = Pt(30)
            with Image.open(chart) as image:
                ratio = image.width / image.height
            width = min(11.0, 4.8 * ratio)
            height = width / ratio
            page.shapes.add_picture(
                str(chart), Inches((13.333 - width) / 2), Inches(1.3),
                width=Inches(width), height=Inches(height),
            )
            summary = next((item.summary for item in result.computations if item.validation_passed and any(path.endswith(chart.name) for path in item.output_files)), "Validated chart")
            caption = page.shapes.add_textbox(Inches(0.9), Inches(6.25), Inches(11.5), Inches(0.5)).text_frame
            caption.text = safe_text(summary)[:180]
            caption.paragraphs[0].font.size = Pt(14)
            footer = page.shapes.add_textbox(Inches(0.9), Inches(7.0), Inches(11.5), Inches(0.3)).text_frame
            footer.text = provenance(result)
            footer.paragraphs[0].font.size = Pt(10)
    elif result.computations:
        slide("Quantitative Analysis", [
            f"{item.computation_id}: {item.summary if item.validation_passed else 'calculation not validated'}"
            for item in result.computations
        ] or ["No validated numerical result"])
    findings = [line.strip() for line in result.final_answer.splitlines() if line.strip() and not line.startswith("Limitations:")][:4]
    slide("Main Results", [*(findings or ["No supported finding was established."]), f"Coverage: {result.goal_coverage_percent:.0f}%"])
    slide("Hypothesis Evaluation", [f"Expected: {result.expected_result or 'Not provided'}", f"Assessment: {result.expected_result_status.value if result.expected_result else 'Not applicable'}"])
    limitations = result.final_limitations or ["No unresolved limitation relevant to the requested goal."]
    for offset in range(0, len(limitations), 5):
        slide("Limitations", limitations[offset:offset + 5])
    slide("Conclusion", [f"Goal: {result.status.value}", f"Stop reason: {result.stop_reason.value if result.stop_reason else 'not recorded'}", f"Iterations: {result.iterations_completed}"])
    deck.save(path)
